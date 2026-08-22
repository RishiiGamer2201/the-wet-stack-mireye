"""Turn the EIA-861 workbooks into one compact JSON keyed by state + county.

    python scripts/build_eia861.py

Run from the repository root. Rebuild this once a year, when EIA publishes the
next Form EIA-861 (usually the following October), by bumping YEAR.

Needs `openpyxl`, which is a *download-time* dependency only: the runtime adapter
reads the JSON this produces and never opens a spreadsheet.

    pip install openpyxl
"""

import io
import json
import pathlib
import sys
import zipfile
from datetime import UTC, datetime

import httpx
import openpyxl

sys.path.insert(0, "apps/api")
from app.adapters.datasets import normalise_county  # noqa: E402

YEAR = 2023
URL = f"https://www.eia.gov/electricity/data/eia861/archive/zip/f861{YEAR}.zip"
OUT = pathlib.Path("apps/api/app/data/datasets/eia861_reliability.json")


def cell(value):
    """EIA writes '.' for 'not reported'. That is not zero and not a number."""
    if value is None:
        return None
    text = str(value).strip()
    if text in {"", ".", "NM", "NA"}:
        return None
    try:
        return round(float(text), 2)
    except ValueError:
        return None


def main() -> None:
    print(f"downloading {URL}")
    response = httpx.get(URL, headers={"User-Agent": "wetstack-mireye/1.0"}, timeout=300, follow_redirects=True)
    response.raise_for_status()
    archive = zipfile.ZipFile(io.BytesIO(response.content))

    # --- reliability: one row per utility per state --------------------------
    wb = openpyxl.load_workbook(io.BytesIO(archive.read(f"Reliability_{YEAR}.xlsx")), read_only=True, data_only=True)
    ws = wb["Reliability_States"]
    rows = ws.iter_rows(min_row=4, values_only=True)
    utilities: dict[str, dict] = {}
    for row in rows:
        if not row or row[1] is None:
            continue
        number = str(row[1]).strip()
        state = (row[3] or "").strip()
        # Columns: 5 SAIDI with MED, 8 SAIDI without MED, 9 SAIFI without MED.
        with_med, without_med, saifi = cell(row[5]), cell(row[8]), cell(row[9])
        if without_med is None and with_med is None:
            continue  # the utility filed the form but reported no reliability data
        utilities[f"{number}|{state}"] = {
            "utility": (row[2] or "").strip(),
            "state": state,
            "ownership": (row[4] or "").strip(),
            "saidi_without_med": without_med,
            "saidi_with_med": with_med,
            "saifi_without_med": saifi,
        }
    wb.close()

    # --- service territory: which utilities serve which county ---------------
    wb = openpyxl.load_workbook(io.BytesIO(archive.read(f"Service_Territory_{YEAR}.xlsx")), read_only=True, data_only=True)
    ws = wb["Counties_States"]
    counties: dict[str, list[str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[1] is None:
            continue
        number, state, county = str(row[1]).strip(), (row[4] or "").strip(), (row[5] or "").strip()
        if not state or not county:
            continue
        key = f"{state}|{normalise_county(county)}"
        entry = f"{number}|{state}"
        if entry in utilities:
            counties.setdefault(key, []).append(entry)
    wb.close()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps(
            {
                "_source": f"EIA Form EIA-861 ({YEAR}): Reliability + Service Territory",
                "_url": URL,
                "_licence": "US Energy Information Administration, public domain (US Government work)",
                "_downloaded_at": datetime.now(UTC).isoformat(),
                "_note": "SAIDI is reported per utility per state. A county is served by "
                "one or more utilities; this file maps county -> utilities, and the "
                "adapter reports each utility separately rather than averaging them.",
                "utilities": utilities,
                "counties": counties,
            },
            separators=(",", ":"),
        ),
        encoding="utf-8",
    )
    reporting = sum(1 for u in utilities.values() if u["saidi_without_med"] is not None)
    print(f"utilities with reliability data: {len(utilities)} ({reporting} report SAIDI without MED)")
    print(f"counties covered: {len(counties)}")
    print(f"wrote {OUT} ({OUT.stat().st_size / 1024:.0f} kB)")


if __name__ == "__main__":
    main()
